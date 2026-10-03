"""AI column-mapping assistant: the LLM proposes where each source column goes in the C360 data product; a person approves.

    python -m trustissues.factory.curate.mapping_assistant prompts                 # build the prompts only (no API call)
    python -m trustissues.factory.curate.mapping_assistant propose --call          # ask the LLM, write the proposals
    python -m trustissues.factory.curate.mapping_assistant propose --call --tables customers card_accounts
    python -m trustissues.factory.curate.mapping_assistant status                  # how many proposed / reviewed
    python -m trustissues.factory.curate.mapping_assistant evaluate                # agreement with the hand-built mappings

Flow (design §7 "Source-to-target mapping", plan §3.2):
  1. Read every source column from DATA_DICTIONARY.xlsx (table, column, type, key, description, example) and its
     sensitivity from config/sensitivity.yaml.
  2. Send the columns of one table, in batches, with the catalogue of C360 targets (TARGETS below) to the LLM.
     It returns, per column: target, transform, confidence (0-1) and a one-line reason, as JSON.
  3. Write data/mappings/column_mappings.csv, lowest confidence first, for the engineer to review. A reviewer sets
     `status` (approved | edited | rejected) and, when editing, `approved_target` / `approved_transform`.
  4. Re-running keeps every reviewed row and only re-asks columns that are new or whose dictionary entry changed.
  5. `evaluate` compares the proposals with REFERENCE: the mappings the pipeline was built with by hand. That is the
     measure of how far the suggestions can be trusted.

Governance:
  - Nothing is sent unless `--call` is given AND the key is set. The key is read from the environment variable named in
    config (`OPENROUTER_API_KEY`) or from the repo's .env file (gitignored); it is never written anywhere.
  - Only metadata leaves the machine: column names, types, descriptions. For direct-PII, protected and proxy columns
    even the dictionary's example value is withheld. No warehouse data is read.
  - AI suggests, a person approves: the CSV is a proposal list. Nothing in the pipeline reads unapproved rows.
  - Responses are cached by a hash of the request, so a rerun never pays twice; every call is logged to
    data/mappings/llm_calls.jsonl (prompt version, model, request hash, columns, outcome).

Provider: OpenRouter's OpenAI-compatible chat API with NVIDIA Nemotron 3 Super (config `mapping_assistant`). When the
team's shared common/llm.py is merged, `_post` should be replaced by it (see merge.md).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path

import yaml

from trustissues.common import config

MAPPINGS_DIR = config.REPO_ROOT / "data" / "mappings"
MAPPINGS_CSV = MAPPINGS_DIR / "column_mappings.csv"
CALL_LOG = MAPPINGS_DIR / "llm_calls.jsonl"
SENSITIVE = {"direct_pii", "protected"}

CSV_FIELDS = ["source_table", "source_column", "source_type", "sensitivity", "proposed_target", "proposed_transform",
              "confidence", "reason", "model", "prompt_version", "dictionary_hash", "status", "approved_target",
              "approved_transform", "reviewer", "reviewed_ts", "note"]
REVIEW_FIELDS = ["status", "approved_target", "approved_transform", "reviewer", "reviewed_ts", "note"]
REVIEWED = {"approved", "edited", "rejected"}


# --------------------------------------------------------------------------- configuration

def settings() -> dict:
    return config.settings()["mapping_assistant"]


def api_key() -> str | None:
    """From the environment, else from the repo's .env (KEY=VALUE lines). Never logged, never written."""
    name = settings()["api_key_env"]
    if os.environ.get(name):
        return os.environ[name]
    env = config.REPO_ROOT / ".env"
    if env.exists():
        for line in env.read_text(encoding="utf-8").splitlines():
            k, sep, v = line.partition("=")
            if sep and k.strip() == name:
                return v.strip().strip('"').strip("'") or None
    return None


# --------------------------------------------------------------------------- targets

# Targets that are not columns of a gold table. The LLM must pick one of TARGETS or one of these.
SPECIAL_TARGETS = {
    "identity.name": "person's name, used to match records across systems (standardised into std_name_*)",
    "identity.dob": "date of birth, used for matching (std_dob)",
    "identity.phone": "phone number, used for matching (std_phone_*_e164)",
    "identity.postal": "postal code, used for matching (std_postal, std_fsa)",
    "identity.employer": "employer or payroll descriptor, weak matching evidence (std_employer)",
    "identity.national_id": "hashed government ID, exact matching key (CRM only)",
    "identity.source_key": "the source system's own customer key (CRM ID, card/loan customer number, 10-digit CIF)",
    "identity.crm_duplicate_pointer": "CRM's own pointer to the record this one duplicates",
    "relationship.joint_holder": "another person who holds the account jointly (never merged)",
    "relationship.co_borrower": "another person on the loan (never merged)",
    "excluded.protected": "protected attribute: never used in decisions or features (kept for fairness testing only)",
    "lineage": "load / audit metadata (batch id, extract time, record hash, source file)",
    "history_only": "kept in curated history and features, not a field of the customer record",
    "not_mapped": "not needed by the C360 product",
}


def _stub_tables() -> dict[str, list[tuple[str, str]]]:
    """{table: [(column, comment)]} from contracts/interfaces/stubs.sql."""
    from trustissues.common.interfaces import STUBS_SQL
    out: dict[str, list[tuple[str, str]]] = {}
    current = None
    for line in STUBS_SQL.read_text(encoding="utf-8").splitlines():
        m = re.match(r"\s*CREATE TABLE IF NOT EXISTS gold\.(\w+)", line)
        if m:
            current = m.group(1)
            out[current] = []
            continue
        if current and line.strip().startswith(");"):
            current = None
            continue
        m = re.match(r"\s*(\w+)\s+[A-Z]+[^,-]*,?\s*(?:--\s*(.*))?$", line)
        if current and m and not line.strip().startswith("--"):
            out[current].append((m.group(1), (m.group(2) or "").strip()))
    return out


TARGET_TABLES = ["c360_customer", "c360_account", "decisioning_customer", "customer_safeguards", "promises_enriched"]


def targets() -> dict[str, str]:
    """target name -> description: gold columns of the C360 product, then the special targets."""
    contract = yaml.safe_load(open(config.REPO_ROOT / "contracts" / "c360_customer.yaml", encoding="utf-8"))
    described = {f["name"]: f.get("description", "") for f in contract["schema"]}
    out: dict[str, str] = {}
    stubs = _stub_tables()
    for t in TARGET_TABLES:
        for col, comment in stubs.get(t, []):
            if t == "c360_customer" or col != "golden_customer_id":
                out[f"{t}.{col}"] = described.get(col, "") if t == "c360_customer" else comment
    return out | SPECIAL_TARGETS


# --------------------------------------------------------------------------- source columns

@dataclass
class SourceColumn:
    table: str
    column: str
    type: str
    key: str
    description: str
    example: str
    sensitivity: str

    @property
    def withheld(self) -> bool:
        return self.sensitivity in SENSITIVE or self.sensitivity == "proxy"

    def payload(self) -> dict:
        """What the LLM is shown. No example value for sensitive columns."""
        d = {"column": self.column, "type": self.type, "description": self.description}
        if self.key:
            d["key"] = self.key
        d["example"] = f"[withheld: {self.sensitivity}]" if self.withheld else self.example
        if self.sensitivity and self.sensitivity != "none":
            d["sensitivity"] = self.sensitivity
        return d

    @property
    def dictionary_hash(self) -> str:
        return hashlib.sha256(json.dumps([self.type, self.key, self.description, self.example, self.sensitivity])
                              .encode()).hexdigest()[:16]


def load_sources(xlsx: Path | None = None, sensitivity: dict | None = None) -> list[SourceColumn]:
    import openpyxl
    xlsx = xlsx or config.maple_data_dir() / "DATA_DICTIONARY.xlsx"
    if sensitivity is None:
        sensitivity = yaml.safe_load(open(config.CONFIG_DIR / "sensitivity.yaml", encoding="utf-8"))
    levels = sensitivity.get("tables") or {}
    proxies = set(sensitivity.get("proxy_columns") or [])
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = list(wb["Columns"].iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        if not d.get("Table") or not d.get("Column"):
            continue
        t, c = str(d["Table"]), str(d["Column"])
        level = "proxy" if f"{t}.{c}" in proxies else (levels.get(t, {}).get(c) or str(d.get("Sensitivity") or "none"))
        out.append(SourceColumn(t, c, str(d.get("Type") or ""), str(d.get("Key") or ""), str(d.get("Description") or ""),
                                str(d.get("Example") or ""), level))
    return out


# --------------------------------------------------------------------------- prompts

PROMPT = """You map source columns of a Canadian bank's collections data into a governed customer data product (C360).
For each source column below, choose exactly one target from the TARGET CATALOGUE, propose the transform needed to get
the source value into the target's form, and say how sure you are.

Rules:
- target must be a catalogue key, written exactly as listed.
- Protected attributes (gender, marital status, citizenship, household, dependants, newcomer, accessibility,
  vulnerability type, accent, child-benefit deposits) map to "excluded.protected", even if they look useful.
- Names, dates of birth, phones, postal codes and employers in product systems are matching evidence: map them to the
  identity.* targets, not straight to c360_customer.
- Monthly history columns (suffix _m01.._m12) and detailed product fields map to "history_only" unless a catalogue
  field holds them exactly.
- transform: short and concrete, e.g. "rename", "cast to DATE", "parse day-first date", "pad to 10 digits",
  "most restrictive value across CRM duplicates", "true if any source is true", "sum across accounts", "none".
- confidence: 0.0-1.0. Use below 0.5 when unsure; a person reviews the lowest first.
- reason: one sentence.

Return JSON only: {{"mappings": [{{"column": "...", "target": "...", "transform": "...", "confidence": 0.0, "reason": "..."}}]}}
with one entry per source column, in the same order.

TARGET CATALOGUE (key: description):
{catalogue}

SOURCE TABLE: {table}
SOURCE COLUMNS:
{columns}"""


@dataclass
class Batch:
    table: str
    columns: list[SourceColumn]
    prompt: str = ""
    request_hash: str = ""


def build_batches(sources: list[SourceColumn], tables: list[str] | None = None, skip: set | None = None) -> list[Batch]:
    s = settings()
    catalogue = "\n".join(f"- {k}: {v}" for k, v in targets().items())
    by_table: dict[str, list[SourceColumn]] = {}
    for c in sources:
        if (tables and c.table not in tables) or (skip and (c.table, c.column) in skip):
            continue
        by_table.setdefault(c.table, []).append(c)
    batches = []
    for table, cols in sorted(by_table.items()):
        for i in range(0, len(cols), s["batch_size"]):
            chunk = cols[i:i + s["batch_size"]]
            prompt = PROMPT.format(catalogue=catalogue, table=table,
                                   columns=json.dumps([c.payload() for c in chunk], ensure_ascii=False, indent=1))
            b = Batch(table, chunk, prompt)
            b.request_hash = hashlib.sha256(json.dumps([s["model"], s["prompt_version"], prompt]).encode()).hexdigest()
            batches.append(b)
    return batches


# --------------------------------------------------------------------------- LLM call (OpenRouter, OpenAI-compatible)

class LLMError(Exception):
    pass


def _post(prompt: str, key: str) -> str:
    """One chat completion. Returns the message text. Retries 429 / 5xx / timeouts with backoff."""
    import httpx
    s = settings()
    body = {"model": s["model"], "temperature": s.get("temperature", 0), "max_tokens": s["max_tokens"],
            "response_format": {"type": "json_object"},
            "messages": [{"role": "user", "content": prompt}], **(s.get("extra_body") or {})}
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json",
               "X-Title": "trustissues mapping assistant"}
    last = None
    for attempt in range(s["max_retries"] + 1):
        try:
            r = httpx.post(f"{s['base_url'].rstrip('/')}/chat/completions", json=body, headers=headers,
                           timeout=s["timeout_secs"])
            if r.status_code == 200:
                msg = r.json()["choices"][0]["message"]
                return msg.get("content") or ""
            if r.status_code in (401, 403):
                raise LLMError(f"HTTP {r.status_code}: key rejected")  # never retried
            last = f"HTTP {r.status_code}: {r.text[:200]}"
            if r.status_code not in (408, 409, 429) and r.status_code < 500:
                break
        except httpx.HTTPError as e:
            last = f"{e.__class__.__name__}: {e}"
        time.sleep(min(2 ** attempt * 2, 30))
    raise LLMError(last or "unknown error")


def parse_response(text: str, batch: Batch, catalogue: dict) -> list[dict]:
    """Validate the model's JSON against the batch. Missing or invalid entries become confidence 0 for review."""
    m = re.search(r"\{.*\}", text or "", re.S)  # tolerate prose or code fences around the JSON
    try:
        items = json.loads(m.group(0))["mappings"] if m else []
    except (json.JSONDecodeError, KeyError, TypeError):
        items = []
    got = {str(i.get("column")): i for i in items if isinstance(i, dict)}
    out = []
    for c in batch.columns:
        i = got.get(c.column)
        if i is None:
            out.append({"column": c.column, "target": "", "transform": "", "confidence": 0.0,
                        "reason": "no answer from the model for this column"})
            continue
        target = str(i.get("target", "")).strip()
        try:
            conf = max(0.0, min(1.0, float(i.get("confidence", 0))))
        except (TypeError, ValueError):
            conf = 0.0
        reason = str(i.get("reason", "")).strip()
        if target not in catalogue:
            reason = f"[target not in catalogue: {target}] " + reason
            conf = 0.0
        if c.sensitivity == "protected" and target != "excluded.protected":
            reason = "[protected column mapped outside excluded.protected] " + reason
            conf = 0.0
        out.append({"column": c.column, "target": target, "transform": str(i.get("transform", "")).strip(),
                    "confidence": round(conf, 3), "reason": reason})
    return out


# --------------------------------------------------------------------------- cache and call log

def _cache_path() -> Path:
    return config.warehouse_path().parent / settings()["cache_file"]


def _cache_get(h: str) -> str | None:
    p = _cache_path()
    if not p.exists():
        return None
    for line in p.read_text(encoding="utf-8").splitlines():
        rec = json.loads(line)
        if rec["hash"] == h:
            return rec["text"]
    return None


def _cache_put(h: str, text: str) -> None:
    p = _cache_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "a", encoding="utf-8") as f:
        f.write(json.dumps({"hash": h, "text": text}, ensure_ascii=False) + "\n")


def _log_call(batch: Batch, outcome: str, secs: float) -> None:
    s = settings()
    MAPPINGS_DIR.mkdir(parents=True, exist_ok=True)
    with open(CALL_LOG, "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": time.strftime("%Y-%m-%d %H:%M:%S"), "model": s["model"],
                            "prompt_version": s["prompt_version"], "request_hash": batch.request_hash,
                            "table": batch.table, "columns": [c.column for c in batch.columns],
                            "outcome": outcome, "secs": round(secs, 1)}) + "\n")


# --------------------------------------------------------------------------- CSV

def read_csv(path: Path = MAPPINGS_CSV) -> dict[tuple[str, str], dict]:
    if not path.exists():
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {(r["source_table"], r["source_column"]): r for r in csv.DictReader(f)}


def write_csv(rows: dict[tuple[str, str], dict], path: Path = MAPPINGS_CSV) -> None:
    """Unreviewed rows first, lowest confidence first; reviewed rows after, by table."""
    def key(r):
        reviewed = r.get("status") in REVIEWED
        conf = float(r["confidence"]) if r.get("confidence") not in (None, "") else 0.0
        return (reviewed, conf if not reviewed else 0.0, r["source_table"], r["source_column"])
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        w.writeheader()
        for r in sorted(rows.values(), key=key):
            w.writerow({k: r.get(k, "") for k in CSV_FIELDS})


# --------------------------------------------------------------------------- the commands

@dataclass
class ProposeResult:
    batches: int = 0
    cached: int = 0
    called: int = 0
    failed: int = 0
    columns: int = 0
    errors: list = field(default_factory=list)


def propose(call: bool, tables: list[str] | None = None, post=None, sources: list[SourceColumn] | None = None,
            path: Path = MAPPINGS_CSV) -> ProposeResult:
    """Ask the LLM for every column that has no reviewed or current proposal. `post(prompt, key) -> text` is
    injectable for tests. Without `call=True` nothing is sent and nothing is written."""
    s = settings()
    sources = sources if sources is not None else load_sources()
    existing = read_csv(path)
    # keep: reviewed rows, and proposals whose dictionary entry has not changed since they were made
    current = {(c.table, c.column): c.dictionary_hash for c in sources}
    skip = {k for k, r in existing.items()
            if r.get("status") in REVIEWED or (r.get("dictionary_hash") == current.get(k) and r.get("proposed_target"))}
    batches = build_batches(sources, tables, skip)
    res = ProposeResult(batches=len(batches))
    if not call:
        return res
    if len(batches) > s["max_batches"]:
        raise LLMError(f"{len(batches)} batches exceeds max_batches={s['max_batches']} (cost guard); pass --tables")
    key = None
    if post is None:
        key = api_key()
        if not key:
            raise LLMError(f"no API key: set {s['api_key_env']} in the environment or in .env")
        post = _post
    catalogue = targets()

    def run(b: Batch):
        t0 = time.time()
        text = _cache_get(b.request_hash)
        if text is not None:
            return b, text, "cached", time.time() - t0
        try:
            text = post(b.prompt, key)
        except LLMError as e:
            _log_call(b, f"error: {e}", time.time() - t0)
            return b, None, f"error: {e}", time.time() - t0
        _cache_put(b.request_hash, text)
        _log_call(b, "ok", time.time() - t0)
        return b, text, "called", time.time() - t0

    with ThreadPoolExecutor(max_workers=s["workers"]) as ex:
        results = list(ex.map(run, batches))
    for b, text, outcome, _ in results:
        if text is None:
            res.failed += 1
            res.errors.append(f"{b.table}: {outcome}")
            continue
        res.cached += outcome == "cached"
        res.called += outcome == "called"
        for c, m in zip(b.columns, parse_response(text, b, catalogue)):
            old = existing.get((c.table, c.column), {})
            existing[(c.table, c.column)] = {
                "source_table": c.table, "source_column": c.column, "source_type": c.type, "sensitivity": c.sensitivity,
                "proposed_target": m["target"], "proposed_transform": m["transform"], "confidence": m["confidence"],
                "reason": m["reason"], "model": s["model"], "prompt_version": s["prompt_version"],
                "dictionary_hash": c.dictionary_hash, "status": "proposed",
                **{k: old.get(k, "") for k in REVIEW_FIELDS if k != "status"}}
            res.columns += 1
    write_csv(existing, path)
    return res


def write_prompts(tables: list[str] | None = None, path: Path | None = None) -> Path:
    """Dry run: the exact prompts that `propose --call` would send, for inspection. No API call."""
    path = path or MAPPINGS_DIR / "prompts_preview.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        for b in build_batches(load_sources(), tables):
            f.write(json.dumps({"table": b.table, "columns": [c.column for c in b.columns],
                                "request_hash": b.request_hash, "prompt": b.prompt}, ensure_ascii=False) + "\n")
    return path


# The mappings the pipeline was actually built with (curate.yaml, identity, build_c360, views). Used only to measure
# the assistant; extend it as more columns are mapped by hand.
REFERENCE = {
    ("customers", "crm_customer_id"): "identity.source_key",
    ("customers", "full_name_raw"): "identity.name",
    ("customers", "dob_raw"): "identity.dob",
    ("customers", "date_of_birth"): "c360_customer.date_of_birth",
    ("customers", "primary_phone_raw"): "identity.phone",
    ("customers", "postal_code_raw"): "identity.postal",
    ("customers", "employer_name"): "identity.employer",
    ("customers", "national_id_hash"): "identity.national_id",
    ("customers", "duplicate_of_crm_id"): "identity.crm_duplicate_pointer",
    ("customers", "segment"): "c360_customer.segment",
    ("customers", "province_code"): "c360_customer.province",
    ("customers", "consent_call"): "c360_customer.consent_phone",
    ("customers", "consent_sms"): "c360_customer.consent_sms",
    ("customers", "consent_email"): "c360_customer.consent_email",
    ("customers", "consent_letter"): "c360_customer.consent_mail",
    ("customers", "cease_communication_flag"): "c360_customer.cease_contact",
    ("customers", "deceased_flag"): "c360_customer.deceased",
    ("customers", "bankruptcy_flag"): "c360_customer.bankruptcy",
    ("customers", "employment_status"): "decisioning_customer.employment_status",
    ("customers", "declared_annual_income"): "decisioning_customer.declared_annual_income",
    ("customers", "housing_status"): "decisioning_customer.housing_status",
    ("customers", "vulnerability_flag"): "customer_safeguards.vulnerability_flag",
    ("customers", "gender_code"): "excluded.protected",
    ("customers", "marital_status"): "excluded.protected",
    ("customers", "citizenship_status"): "excluded.protected",
    ("card_accounts", "src_customer_ref"): "identity.source_key",
    ("card_accounts", "cardholder_name_raw"): "identity.name",
    ("card_accounts", "cardholder_dob_raw"): "identity.dob",
    ("card_accounts", "cardholder_phone_raw"): "identity.phone",
    ("card_accounts", "card_account_id"): "c360_account.account_id",
    ("card_accounts", "current_balance"): "c360_account.balance",
    ("card_accounts", "past_due_amount"): "c360_account.overdue_amount",
    ("card_accounts", "dpd"): "c360_account.dpd",
    ("card_accounts", "dpd_bucket"): "c360_account.bucket",
    ("card_accounts", "account_status"): "c360_account.status",
    ("card_accounts", "dispute_open_flag"): "c360_customer.open_dispute",
    ("loan_accounts", "borrower_name_raw"): "identity.name",
    ("loan_accounts", "co_borrower_ref"): "relationship.co_borrower",
    ("loan_accounts", "total_outstanding"): "c360_account.balance",
    ("deposit_accounts", "src_customer_ref"): "identity.source_key",
    ("deposit_accounts", "holder_name_raw"): "identity.name",
    ("deposit_accounts", "joint_holder_ref"): "relationship.joint_holder",
    ("deposit_accounts", "ccb_deposit_flag"): "excluded.protected",
    ("external", "bank_subject_ref"): "identity.source_key",
    ("external", "subject_name_reported"): "identity.name",
    ("external", "deceased_indicator"): "c360_customer.deceased",
    ("external", "bankruptcy_on_file_flag"): "c360_customer.bankruptcy",
    ("external", "risk_score"): "decisioning_customer.bureau_score",
    ("external", "risk_score_model"): "decisioning_customer.bureau_score_model",
    ("collections_cases", "cease_contact_flag"): "c360_customer.cease_contact",
    ("collections_cases", "hardship_flag"): "c360_customer.hardship_flag",
    ("promises_to_pay", "ptp_status"): "promises_enriched.ptp_status",
    ("promises_to_pay", "ptp_created_ts"): "promises_enriched.ptp_created_ts",
}


def evaluate(path: Path = MAPPINGS_CSV) -> dict:
    """How often the assistant's first proposal equals the hand-built mapping."""
    rows = read_csv(path)
    scored = [(k, rows[k].get("proposed_target"), want) for k, want in REFERENCE.items() if k in rows]
    agree = [k for k, got, want in scored if got == want]
    return {"reference": len(REFERENCE), "proposed": len(scored), "agree": len(agree),
            "agreement": round(len(agree) / len(scored), 3) if scored else None,
            "disagreements": [(f"{t}.{c}", got, want) for (t, c), got, want in scored if got != want]}


def status(path: Path = MAPPINGS_CSV) -> dict:
    rows = read_csv(path).values()
    out: dict = {}
    for r in rows:
        out[r.get("status") or "proposed"] = out.get(r.get("status") or "proposed", 0) + 1
    out["low_confidence_unreviewed"] = sum(1 for r in rows if r.get("status") not in REVIEWED
                                           and float(r.get("confidence") or 0) < 0.5)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="AI column-mapping assistant (proposals for a person to approve)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("prompts", help="write the prompts that would be sent; no API call")
    p.add_argument("--tables", nargs="*")
    p = sub.add_parser("propose", help="ask the LLM; needs --call")
    p.add_argument("--tables", nargs="*")
    p.add_argument("--call", action="store_true", help="actually send requests to the API")
    sub.add_parser("status")
    sub.add_parser("evaluate")
    a = ap.parse_args()
    if a.cmd == "prompts":
        print(f"wrote {write_prompts(a.tables)}")
    elif a.cmd == "propose":
        r = propose(a.call, a.tables)
        if not a.call:
            print(f"{r.batches} batch(es) would be sent. Nothing was sent: add --call to send them.")
        else:
            print(r)
    elif a.cmd == "status":
        print(status())
    else:
        print(json.dumps(evaluate(), indent=1))


if __name__ == "__main__":
    main()
